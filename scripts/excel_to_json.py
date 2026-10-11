"""Export the supplied workbook without inventing missing values."""
import json
from pathlib import Path
from openpyxl import load_workbook

root = Path(__file__).resolve().parents[1]
sheet = load_workbook(root / 'home_medicine_inventory.xlsx', data_only=True)['Лекарства']
rows = []
for cells in sheet.iter_rows(min_row=7, values_only=True):
    if not isinstance(cells[0], int):
        continue
    number, name, dosage, form, quantity, unit, expiry, precision, location, source, verify, status, comment = cells
    quantity_is_unknown = isinstance(comment, str) and 'Количество' in comment and 'не указан' in comment
    rows.append(dict(source_id=number, name=name, dosage=dosage, form=form,
                     quantity=None if quantity_is_unknown else quantity, unit=unit, expires_on=expiry,
                     expiry_precision=precision, storage_location=location,
                     source_text=source, needs_review=verify == 'Да',
                     source_status=status, review_note=comment))
assert len(rows) == 95
assert sum(row['needs_review'] for row in rows) == 43
assert {row['storage_location'] for row in rows} == {'Зелёный ящик', 'Нижний ящик'}
target = root / 'src' / 'data' / 'inventory.json'
target.parent.mkdir(parents=True, exist_ok=True)
target.write_text(json.dumps(rows, ensure_ascii=False, indent=2) + '\n', encoding='utf-8')
print(f'Exported {len(rows)} records to {target}')
